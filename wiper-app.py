
# app.py
import os
from typing import Dict, List, Optional

import pandas as pd
import streamlit as st

st.set_page_config(page_title="Glove Finder", page_icon="🧤", layout="wide")
st.title("Glove Finder")
st.caption("Find the right glove by cut level/category, colour and safety attributes.")

# =========================================================
# 🔧 DISPLAY ORDER — adjust here (app creator only)
# Valid labels: "Colour", "Cut Category", "Cut rating",
#               "Food Safe?", "Chemical rated?", "Heat rated?"
ORDER_LEFT  = ["Cut Category - (new style letter type)", "Cut rating (old style number type)", "Colour (predominant colour only)"]
ORDER_RIGHT: List[str] = ["Food Safe?", "Chemical rated?", "Heat rated?"]
# Set to True to show the raw option lists near the controls for diagnostics
DEBUG = False
# =========================================================


# -----------------------------
# Embedded image extraction
# -----------------------------
@st.cache_data(show_spinner=False)
def extract_embedded_images(xlsx_path: str, images_dir: str = "images") -> Dict[int, str]:
    """
    Extract embedded images from the first worksheet and map Excel row -> saved image path.
    Returns: {excel_row_number: image_file_path}
    """
    mapping: Dict[int, str] = {}
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.cell import coordinate_to_tuple
    except Exception:
        return mapping

    try:
        wb = load_workbook(xlsx_path, data_only=True)
        ws = wb.active
        os.makedirs(images_dir, exist_ok=True)
        imgs = getattr(ws, "_images", [])
        for idx, img in enumerate(imgs, start=1):
            anchor = getattr(img, "anchor", None)
            row = None
            # Derive row from anchor when possible
            try:
                if isinstance(anchor, str):
                    row = coordinate_to_tuple(anchor)[0]  # (row, col)
                else:
                    from_anchor = getattr(anchor, "_from", None)
                    if from_anchor is not None:
                        row = int(from_anchor.row) + 1
            except Exception:
                row = None

            filename = f"row_{row or idx}.png"
            out_path = os.path.join(images_dir, filename)
            try:
                data = img._data()  # type: ignore[attr-defined]
                with open(out_path, "wb") as f:
                    f.write(data)
                if row:
                    mapping[row] = out_path
            except Exception:
                # Fallback via PIL if available
                try:
                    from PIL import Image as PILImage
                    pil = getattr(img, "_image", None)
                    if pil is not None and isinstance(pil, PILImage.Image):
                        pil.save(out_path)
                        if row:
                            mapping[row] = out_path
                except Exception:
                    pass
    except Exception:
        return mapping
    return mapping


# -----------------------------
# Utilities
# -----------------------------
PLACEHOLDERS = {"", "-", "n/a", "na", "none", "null", "not en 388 rated"}

def _norm_str(x: object) -> str:
    """Normalise for comparisons: casefolded + stripped; NaN -> empty."""
    if pd.isna(x):
        return ""
    return str(x).strip().casefold()

def _display_str(x: object) -> str:
    """Clean display string (preserve case, strip only)."""
    if pd.isna(x):
        return ""
    return str(x).strip()

def _as_int_string_if_number(s: str) -> str:
    """Turn '1.0' -> '1' for nicer Cut values; leave non-numeric (e.g., 'X') untouched."""
    try:
        f = float(s)
        if f.is_integer():
            return str(int(f))
        return s
    except Exception:
        return s

def _resolve_column(preferred_aliases: List[str], df_cols_lower: Dict[str, str]) -> Optional[str]:
    """
    Given a list of lowercased aliases (priority order), return the actual DataFrame column
    name that matches, else None.
    """
    for alias in preferred_aliases:
        if alias in df_cols_lower:
            return df_cols_lower[alias]
    return None


# -----------------------------
# Data loader
# -----------------------------
@st.cache_data(show_spinner=False)
def load_data(path: str) -> pd.DataFrame:
    row_to_img = extract_embedded_images(path)

    df = pd.read_excel(path, engine="openpyxl")
    df.columns = [c.strip() for c in df.columns]

    # Build lowercase->original-name mapping to support aliases
    df_cols_lower = {c.lower(): c for c in df.columns}

    # Resolve the three dropdown columns (case/alias tolerant)
    colour_col = _resolve_column(["colour", "color"], df_cols_lower) or "Colour"
    cutcat_col = _resolve_column(["cut category", "cutcategory"], df_cols_lower) or "Cut Category"
    cut_col    = _resolve_column(["cut"], df_cols_lower) or "Cut"

    # Warn if any target columns truly missing
    missing = [col for col in [colour_col, cutcat_col, cut_col] if col not in df.columns]
    if missing:
        st.error(f"Missing expected column(s): {missing}. Please check workbook headers.")

    # Ensure Image column; fill from extracted mapping by Excel row index (header=1)
    if "Image" not in df.columns:
        df["Image"] = None
    for i in range(len(df)):
        excel_row = i + 2
        if pd.isna(df.at[i, "Image"]) or not str(df.at[i, "Image"]).strip():
            if excel_row in row_to_img:
                df.at[i, "Image"] = row_to_img[excel_row]

    # Boolean normalisation for Yes-only filters
    def to_bool(x):
        s = _norm_str(x)
        if s in {"yes", "y", "true", "1"}:
            return True
        if s in {"no", "n", "false", "0", "-"}:
            return False
        return None

    for col in ["Food Safe", "Chemical Resistance", "Heat Resistance"]:
        if col in df.columns:
            df[col + " (bool)"] = df[col].apply(to_bool)

    # Create normalised copies for reliable filtering (keep originals for display)
    if colour_col in df.columns:
        df["_colour_norm"] = df[colour_col].astype(str).str.strip().str.casefold()
    if cutcat_col in df.columns:
        df["_cutcat_norm"] = df[cutcat_col].astype(str).str.strip().str.casefold()
    if cut_col in df.columns:
        # Make a clean display variant for Cut (1.0 -> 1)
        df["_cut_display"] = df[cut_col].astype(str).str.strip().map(_as_int_string_if_number)
        df["_cut_norm"] = df["_cut_display"].str.strip().str.casefold()

    # Save resolved names for global use
    df.attrs["colour_col"] = colour_col
    df.attrs["cutcat_col"] = cutcat_col
    df.attrs["cut_col"] = cut_col

    return df


DATA_PATH = "wurth_safety_glove_MASTER List.xlsx"
df = load_data(DATA_PATH)

# Resolved columns (for option builders and display)
COLOUR_COL = df.attrs.get("colour_col", "Colour")
CUTCAT_COL = df.attrs.get("cutcat_col", "Cut Category")
CUT_COL    = df.attrs.get("cut_col", "Cut")


# -----------------------------
# Option builders (robust)
# -----------------------------
def _clean_distinct(series: pd.Series) -> List[str]:
    """
    Distinct, non-placeholder strings for dropdowns, preserving display case.
    If series is empty, returns [].
    """
    if series is None or series.empty:
        return []
    series = series.dropna()
    vals: List[str] = []
    for v in series:
        disp = _display_str(v)
        norm = _norm_str(v)
        if not disp:
            continue
        if norm in PLACEHOLDERS:
            continue
        vals.append(disp)
    # Deduplicate by the normalised value, preserve first display string
    seen = set()
    uniq: List[str] = []
    for v in vals:
        k = _norm_str(v)
        if k not in seen:
            seen.add(k)
            uniq.append(v)
    return uniq

def _fallback_distinct(series: pd.Series) -> List[str]:
    """
    Fallback that only trims and drops NaN; used if cleaned result is empty.
    Guarantees we still show something from the underlying column.
    """
    if series is None or series.empty:
        return []
    series = series.dropna().astype(str).str.strip()
    uniq = [v for v in series.unique().tolist() if v]
    return uniq

def options_for(label: str) -> List[str]:
    """
    Build option list for non-boolean filters (first option 'Any').
    Hardened against placeholders/blanks and with fallbacks so lists never collapse.
    """
    if label == "Colour" and COLOUR_COL in df.columns:
        uniq = _clean_distinct(df[COLOUR_COL])
        if not uniq:
            uniq = _fallback_distinct(df[COLOUR_COL])
        uniq.sort(key=lambda x: x.casefold())
        return ["Any"] + uniq

    if label == "Cut Category" and CUTCAT_COL in df.columns:
        uniq = _clean_distinct(df[CUTCAT_COL])
        if not uniq:
            uniq = _fallback_distinct(df[CUTCAT_COL])
        # Preferred order A..F; 'X'/others after
        order_map = {k: i for i, k in enumerate(list("ABCDEF"))}
        uniq.sort(key=lambda x: (order_map.get(x.upper(), 99), x.upper()))
        return ["Any"] + uniq

    if label == "Cut rating":
        if "_cut_display" in df.columns:
            uniq = _clean_distinct(df["_cut_display"])
            if not uniq:
                uniq = _fallback_distinct(df["_cut_display"])
        elif CUT_COL in df.columns:
            base = df[CUT_COL].astype(str).str.strip().map(_as_int_string_if_number)
            uniq = _clean_distinct(base)
            if not uniq:
                uniq = _fallback_distinct(base)
        else:
            uniq = []
        # Numbers first (1,2,3...), then tokens like 'X'
        def _key(x: str):
            return (0, int(x)) if x.isdigit() else (1, x.upper())
        uniq.sort(key=_key)
        return ["Any"] + uniq

    return ["Any"]


# -----------------------------
# Filter UI (two columns, code-ordered)
# -----------------------------
with st.container():
    st.subheader("Filter")

    left_col, right_col = st.columns(2)
    selections: Dict[str, object] = {}

    def render_filter(label: str, col):
        # Yes-only checkboxes
        if label in {"Food Safe?", "Chemical rated?", "Heat rated?"}:
            selections[label] = col.checkbox(label + " (Yes only)", value=False, key=f"cb_{label}")
            return

        # Dropdowns
        opts = options_for(label)
        if DEBUG:
            col.caption(f"{label} options: {opts}")
        selections[label] = col.selectbox(label, opts, index=0, key=f"sb_{label}")

    for lbl in ORDER_LEFT:
        render_filter(lbl, left_col)
    for lbl in ORDER_RIGHT:
        render_filter(lbl, right_col)

    go = st.button("Search", type="primary")


# -----------------------------
# Apply filters
# -----------------------------
if go:
    filtered = df.copy()

    # Dropdown filters (case-insensitive, trimmed)
    if "Colour" in selections and isinstance(selections["Colour"], str) and selections["Colour"] != "Any":
        right = _norm_str(selections["Colour"])
        if "_colour_norm" in filtered.columns:
            filtered = filtered[filtered["_colour_norm"] == right]

    if "Cut Category" in selections and isinstance(selections["Cut Category"], str) and selections["Cut Category"] != "Any":
        right = _norm_str(selections["Cut Category"])
        if "_cutcat_norm" in filtered.columns:
            filtered = filtered[filtered["_cutcat_norm"] == right]

    if "Cut rating" in selections and isinstance(selections["Cut rating"], str) and selections["Cut rating"] != "Any":
        right = _norm_str(_as_int_string_if_number(selections["Cut rating"]))
        if "_cut_norm" in filtered.columns:
            filtered = filtered[filtered["_cut_norm"] == right]

    # Yes-only boolean filters
    def apply_yes_only(local_df: pd.DataFrame, src_label: str, sel_key: str) -> pd.DataFrame:
        yes_checked = bool(selections.get(sel_key, False))
        col_bool = src_label + " (bool)"
        if yes_checked and col_bool in local_df.columns:
            mask = local_df[col_bool].fillna(False) == True  # element-wise
            return local_df[mask]
        return local_df

    filtered = apply_yes_only(filtered, "Food Safe", "Food Safe?")
    filtered = apply_yes_only(filtered, "Chemical Resistance", "Chemical rated?")
    filtered = apply_yes_only(filtered, "Heat Resistance", "Heat rated?")

    st.subheader("Search results")
    st.write(f"{len(filtered)} result(s)")

    # ---- Results cards ----
    for _, row in filtered.iterrows():
        with st.container(border=True):
            left, right = st.columns([1, 3])

            # Image
            img_src = row.get("Image", None)
            if isinstance(img_src, str) and img_src.strip():
                try:
                    left.image(img_src, use_column_width=True)
                except Exception:
                    left.info("No image available")
            else:
                left.info("No image available")

            # Name + Product link
            name = row.get("Glove Name", "(no name)")
            right.markdown(f"### {name}")
            link = row.get("Product Link", "")
            if isinstance(link, str) and link.strip():
                right.link_button("View product", link)

            # Attributes grid (compact) with integer formatting for EN388 sub-ratings
            attrs: List[tuple] = []
            for label in [
                "Article Numbers", "Colour", "EN 388 Code", "Abrasion", "Cut", "Tear", "Puncture",
                "Cut Category", "Impact", "Chemical Resistance", "Heat Resistance", "Food Safe", "Tactile"
            ]:
                val = row.get(label, None)
                if pd.isna(val):
                    continue
                # Show EN388 sub-ratings with no decimals
                if label in ["Abrasion", "Cut", "Tear", "Puncture"] and isinstance(val, (int, float)):
                    val = int(val)
                attrs.append((label, str(val)))

            a1, a2 = right.columns(2)
            half = (len(attrs) + 1) // 2
            for col, items in [(a1, attrs[:half]), (a2, attrs[half:])]:
                for lab, v in items:
                    col.markdown(f"**{lab}:** {v}")
            st.divider()

    if not filtered.empty:
        csv = filtered.to_csv(index=False)
        st.download_button(
            "Download results (CSV)",
            data=csv,
            file_name="glove_finder_results.csv",
            mime="text/csv",
        )
else:
    st.info("Choose filters above and press **Search** to see matching gloves.")
